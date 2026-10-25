# -*- coding: utf-8 -*-
"""Dump full structure of a DOCX: paragraphs, runs, tables, sections, headers/footers."""
import sys, os, json
from docx import Document
from docx.shared import Pt

path = sys.argv[1]
out = sys.argv[2]

doc = Document(path)

lines = []
lines.append(f"FILE: {path}")
lines.append(f"size: {os.path.getsize(path)} bytes")

# sections
lines.append("\n===== SECTIONS =====")
for i, s in enumerate(doc.sections):
    lines.append(
        f"[{i}] page={s.page_width.cm:.2f}x{s.page_height.cm:.2f}cm "
        f"margins L/R/T/B={s.left_margin.cm:.2f}/{s.right_margin.cm:.2f}/{s.top_margin.cm:.2f}/{s.bottom_margin.cm:.2f}cm "
        f"header_dist={s.header_distance.cm:.2f} footer_dist={s.footer_distance.cm:.2f}"
    )
    try:
        hp = " | ".join(p.text for p in s.header.paragraphs if p.text.strip())
        fp = " | ".join(p.text for p in s.footer.paragraphs if p.text.strip())
        lines.append(f"    header: {hp!r}")
        lines.append(f"    footer: {fp!r}")
    except Exception as e:
        lines.append(f"    header/footer error: {e}")

# body: iterate in document order
lines.append("\n===== BODY (document order) =====")
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

def iter_block_items(parent):
    from docx.document import Document as _Doc
    if isinstance(parent, _Doc):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

pi = 0
ti = 0
for block in iter_block_items(doc):
    if isinstance(block, Paragraph):
        txt = block.text
        # collect run fonts
        fonts = []
        for r in block.runs:
            if r.text.strip():
                nm = r.font.name
                ea = None
                try:
                    rpr = r._element.find(qn('w:rPr'))
                    if rpr is not None:
                        rf = rpr.find(qn('w:rFonts'))
                        if rf is not None:
                            ea = rf.get(qn('w:eastAsia'))
                except Exception:
                    pass
                sz = r.font.size.pt if r.font.size else None
                b = r.font.bold
                fonts.append(f"{nm}/{ea}/{sz}/B={b}")
        align = block.alignment
        pf = block.paragraph_format
        spacing = f"before={pf.space_before.pt if pf.space_before else None},after={pf.space_after.pt if pf.space_after else None},line={pf.line_spacing}"
        indent = f"first={pf.first_line_indent.cm if pf.first_line_indent else None},left={pf.left_indent.cm if pf.left_indent else None}"
        marker = "P" if txt.strip() else "-"
        lines.append(f"[{marker}{pi:03d}] style={block.style.name!r} align={align} {spacing} {indent}")
        lines.append(f"      TEXT: {txt}")
        if fonts:
            lines.append(f"      RUNS: {' ; '.join(fonts[:8])}")
        pi += 1
    else:
        lines.append(f"\n[TABLE {ti}] style={block.style.name if block.style else None} rows={len(block.rows)} cols={len(block.columns)}")
        for ri, row in enumerate(block.rows):
            cells = []
            seen = set()
            for ci, cell in enumerate(row.cells):
                key = id(cell._tc)
                if key in seen:
                    cells.append("<merged>")
                    continue
                seen.add(key)
                cells.append(cell.text.replace("\n", " ⏎ "))
            lines.append(f"   R{ri}: " + " || ".join(cells))
        ti += 1

lines.append(f"\nTOTAL paragraphs={pi} tables={ti}")
lines.append(f"inline_shapes={len(doc.inline_shapes)}")
for i, sh in enumerate(doc.inline_shapes):
    lines.append(f"  shape[{i}] type={sh.type} w={sh.width.cm:.2f}cm h={sh.height.cm:.2f}cm")

# styles used
lines.append("\n===== STYLES USED =====")
used = {}
for p in doc.paragraphs:
    used.setdefault(p.style.name, 0)
    used[p.style.name] += 1
for k, v in sorted(used.items(), key=lambda x: -x[1]):
    lines.append(f"  {k}: {v}")

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"written {out}, lines={len(lines)}, paragraphs={pi}, tables={ti}")
